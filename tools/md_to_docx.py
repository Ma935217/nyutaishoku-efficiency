# -*- coding: utf-8 -*-
"""Markdown（見出し・箇条書き・番号付き・表・引用・太字）を水色基調のWord文書に変換する。
使い方: python tools/md_to_docx.py 入力.md [出力.docx]
対応は本リポジトリの文書で使う範囲に限定したシンプルな変換。"""
import re, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE_DARK = RGBColor(0x1F, 0x5F, 0x8B)
BLUE_MID = RGBColor(0x2E, 0x86, 0xC1)
HEADER_FILL = "D6EAF8"   # 表見出しの水色
QUOTE_FILL = "EBF5FB"
FONT_JP = "游ゴシック"
FONT_EN = "Calibri"


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT_EN
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT_JP)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color is not None: run.font.color.rgb = color


def shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def add_inline(par, text, size=10.5, color=None):
    """**太字** と `コード` を解釈して run を追加する。"""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part: continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); set_font(r, size, True, color)
        elif part.startswith("`"):
            r = par.add_run(part[1:-1]); set_font(r, size, None, BLUE_DARK)
        else:
            r = par.add_run(part); set_font(r, size, None, color)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for i, r in enumerate(cells) if not (i == 1 and all(set(c) <= set("-: ") for c in r))]
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Table Grid"
    for i, r in enumerate(cells):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            txt = r[j] if j < len(r) else ""
            add_inline(p, txt, 9, BLUE_DARK if i == 0 else None)
            if i == 0:
                shade(cell, HEADER_FILL)
                for run in p.runs: run.font.bold = True
    doc.add_paragraph()


def convert(md_path, out_path):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]; st.font.size = Pt(10.5)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:eastAsia"), FONT_JP)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            add_table(doc, block); continue
        m = re.match(r"^(#{1,3})\s+(.*)", ln)
        if m:
            lvl = len(m.group(1)); text = m.group(2)
            p = doc.add_paragraph()
            if lvl == 1:
                add_inline(p, text, 18, BLUE_DARK); p.runs[0].font.bold = True
                pb = p.paragraph_format; pb.space_after = Pt(6)
                # 下線代わりの水色罫線
                ppr = p._element.get_or_add_pPr(); bdr = OxmlElement("w:pBdr")
                b = OxmlElement("w:bottom"); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "12"); b.set(qn("w:color"), "2E86C1")
                bdr.append(b); ppr.append(bdr)
            elif lvl == 2:
                add_inline(p, text, 14, BLUE_MID); p.runs[0].font.bold = True
                p.paragraph_format.space_before = Pt(12)
            else:
                add_inline(p, text, 12, BLUE_DARK); p.runs[0].font.bold = True
                p.paragraph_format.space_before = Pt(8)
            i += 1; continue
        if ln.startswith(">"):
            p = doc.add_paragraph(); add_inline(p, ln.lstrip("> ").strip(), 9.5)
            p.paragraph_format.left_indent = Cm(0.5)
            ppr = p._element.get_or_add_pPr(); shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), QUOTE_FILL); ppr.append(shd)
            i += 1; continue
        if ln.strip() == "---":
            i += 1; continue
        m = re.match(r"^(\s*)[-*]\s+(.*)", ln)
        if m:
            depth = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet" if depth == 0 else "List Bullet 2")
            add_inline(p, m.group(2)); i += 1; continue
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)", ln)
        if m:
            p = doc.add_paragraph(); add_inline(p, m.group(2) + ". " + m.group(3))
            p.paragraph_format.left_indent = Cm(0.8 + 0.5 * (len(m.group(1)) // 2))
            p.paragraph_format.first_line_indent = Cm(-0.6)
            i += 1; continue
        if ln.strip().startswith("- [ ]"):
            p = doc.add_paragraph(); add_inline(p, "☐ " + ln.strip()[5:].strip()); i += 1; continue
        if ln.strip() == "":
            i += 1; continue
        p = doc.add_paragraph(); add_inline(p, ln.strip()); i += 1
    doc.save(out_path)


if __name__ == "__main__":
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    convert(src, dst)
    print("書き出し:", dst)
